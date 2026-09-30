import streamlit as st
import pandas as pd
import plotly.express as px
from io import BytesIO
import re
import html
import matplotlib.pyplot as plt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.units import cm


# =========================================================
# KONFIGURASI
# =========================================================

st.set_page_config(
    page_title="Dashboard Produktivitas 2026",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)




# =========================================================
# DESAIN DASHBOARD
# =========================================================

st.markdown("""
<style>
/* =========================================================
   PREMIUM DASHBOARD — LIGHT / GLASS / MODERN
   Terinspirasi dari referensi yang diberikan:
   pastel blue, soft glass cards, rounded corners, clean spacing.
   ========================================================= */

:root {
    --navy: #12324a;
    --teal: #0b8f91;
    --teal-dark: #066d73;
    --blue: #237fc5;
    --sky: #dff1f7;
    --ink: #19354a;
    --muted: #60788b;
    --line: #c7dce5;
    --card: rgba(255,255,255,.94);
    --shadow: 0 14px 35px rgba(38, 78, 103, .10);
}

/* ========================= APP CANVAS ========================= */
[data-testid="stAppViewContainer"] {
    background:
        radial-gradient(circle at 88% 4%, rgba(35, 127, 197, .26), transparent 28%),
        radial-gradient(circle at 4% 78%, rgba(11, 143, 145, .18), transparent 24%),
        linear-gradient(135deg, #e3f3f7 0%, #f0f5f8 48%, #e5f0fa 100%);
}

[data-testid="stHeader"] {
    background: rgba(255,255,255,.55);
    border-bottom: 1px solid rgba(210,226,235,.65);
    backdrop-filter: blur(14px);
}

.block-container {
    padding-top: 1.1rem;
    padding-bottom: 2.8rem;
    max-width: 1480px;
}

/* Hide Streamlit's extra chrome where possible */
#MainMenu { visibility: hidden; }
footer { visibility: hidden; }

/* ========================= TYPOGRAPHY ========================= */
h1, h2, h3, h4 {
    color: var(--navy) !important;
    letter-spacing: -.025em;
}
p, li {
    color: #4d6578;
}
[data-testid="stCaptionContainer"] {
    color: var(--muted) !important;
}

/* ========================= SIDEBAR ========================= */
section[data-testid="stSidebar"] {
    background:
        linear-gradient(180deg, #f7fbfc 0%, #d9ecef 100%);
    border-right: 1px solid #d7e7eb;
    box-shadow: 8px 0 28px rgba(44, 87, 109, .06);
}

section[data-testid="stSidebar"] > div {
    background: transparent;
}

section[data-testid="stSidebar"] h1,
section[data-testid="stSidebar"] h2,
section[data-testid="stSidebar"] h3,
section[data-testid="stSidebar"] p,
section[data-testid="stSidebar"] span,
section[data-testid="stSidebar"] label,
section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p {
    color: var(--ink) !important;
    -webkit-text-fill-color: var(--ink) !important;
    opacity: 1 !important;
}

section[data-testid="stSidebar"] [data-testid="stRadio"] > label {
    color: #7890a1 !important;
    font-size: 11px !important;
    font-weight: 700 !important;
    text-transform: uppercase;
    letter-spacing: .08em;
    margin-bottom: 7px;
}

section[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] {
    gap: 5px;
}

section[data-testid="stSidebar"] [data-testid="stRadio"] label {
    border-radius: 12px;
    padding: 8px 10px !important;
    margin: 2px 0 !important;
    transition: .18s ease;
}

section[data-testid="stSidebar"] [data-testid="stRadio"] label:hover {
    background: #d7eeee;
}

section[data-testid="stSidebar"] [data-testid="stRadio"] label[data-checked="true"],
section[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked) {
    background: linear-gradient(135deg, #087d82, #0fa6a6);
    box-shadow: 0 7px 16px rgba(21, 156, 156, .22);
}

section[data-testid="stSidebar"] [data-testid="stRadio"] label[data-checked="true"] *,
section[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked) * {
    color: #fff !important;
    -webkit-text-fill-color: #fff !important;
}

section[data-testid="stSidebar"] hr {
    border-color: #dce9ed !important;
}

.sidebar-brand {
    padding: 8px 3px 18px;
}
.sidebar-brand .brand-mark {
    width: 42px;
    height: 42px;
    border-radius: 13px;
    display: inline-flex;
    align-items: center;
    justify-content: center;
    background: linear-gradient(135deg,#087d82,#237fc5);
    color: #fff;
    font-size: 20px;
    box-shadow: 0 8px 20px rgba(21,156,156,.20);
    margin-bottom: 9px;
}
.sidebar-brand .brand-title {
    font-size: 18px;
    font-weight: 800;
    color: var(--navy);
    line-height: 1.1;
}
.sidebar-brand .brand-sub {
    font-size: 11px;
    color: #8297a7;
    margin-top: 4px;
}

.sidebar-file {
    background: rgba(255,255,255,.75);
    border: 1px solid #d8e8ed;
    border-radius: 12px;
    padding: 10px 12px;
    margin: 8px 0 14px;
    font-size: 11px;
    color: #5f7485;
}
.sidebar-file strong { color: #193b52; }

/* ========================= SELECTBOX ========================= */
section[data-testid="stSidebar"] [data-testid="stSelectbox"] {
    width: 100% !important;
}
section[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] {
    width: 100% !important;
}
section[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] > div {
    background: rgba(255,255,255,.82) !important;
    border: 1px solid #cfe1e7 !important;
    border-radius: 11px !important;
    min-height: 40px !important;
    box-shadow: none !important;
}
section[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] div,
section[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] span,
section[data-testid="stSidebar"] [data-testid="stSelectbox"] [data-baseweb="select"] p {
    color: #34546a !important;
    -webkit-text-fill-color: #34546a !important;
    opacity: 1 !important;
}
section[data-testid="stSidebar"] [data-testid="stSelectbox"] svg {
    color: #159c9c !important;
    fill: #159c9c !important;
}

/* ========================= HERO ========================= */
.dashboard-hero {
    position: relative;
    overflow: hidden;
    background:
        radial-gradient(circle at 94% 20%, rgba(255,255,255,.58), transparent 20%),
        linear-gradient(135deg, #bfe7e6 0%, #d2ebf3 48%, #dfe7f7 100%);
    border: 1px solid rgba(188, 216, 225, .85);
    border-radius: 24px;
    padding: 28px 32px;
    margin-bottom: 22px;
    box-shadow: 0 18px 42px rgba(52, 92, 113, .10);
}
.dashboard-hero:after {
    content: "";
    position: absolute;
    width: 180px;
    height: 180px;
    border-radius: 50%;
    right: -65px;
    top: -75px;
    background: rgba(54, 169, 206, .10);
}
.dashboard-hero .eyebrow {
    display: inline-block;
    background: rgba(255,255,255,.84);
    border: 1px solid rgba(190,218,226,.8);
    color: #087d82;
    border-radius: 999px;
    padding: 6px 11px;
    font-size: 10px;
    font-weight: 800;
    letter-spacing: .10em;
    text-transform: uppercase;
    margin-bottom: 9px;
}
.dashboard-hero h1 {
    color: #17324d !important;
    font-size: 31px;
    line-height: 1.05;
    margin: 0 0 8px;
}
.dashboard-hero p {
    max-width: 760px;
    color: #587183 !important;
    font-size: 14px;
    margin: 0;
}

/* ========================= OVERVIEW HIGHLIGHTS ========================= */
.overview-highlights {
    display: grid;
    grid-template-columns: repeat(3, minmax(0, 1fr));
    gap: 12px;
    margin: -2px 0 22px;
}
.overview-highlight {
    min-height: 86px;
    display: flex;
    align-items: center;
    gap: 12px;
    padding: 14px 16px;
    border-radius: 16px;
    color: #fff;
    box-shadow: 0 10px 24px rgba(38,78,103,.12);
}
.overview-highlight-gold {
    background: linear-gradient(135deg, #a97912 0%, #d2a43a 100%);
}
.overview-highlight-teal {
    background: linear-gradient(135deg, #087d82 0%, #159c9c 100%);
}
.overview-highlight-blue {
    background: linear-gradient(135deg, #176eaf 0%, #237fc5 100%);
}
.overview-highlight-icon {
    width: 42px;
    height: 42px;
    flex: 0 0 42px;
    display: flex;
    align-items: center;
    justify-content: center;
    border-radius: 12px;
    background: rgba(255,255,255,.18);
    border: 1px solid rgba(255,255,255,.24);
    font-size: 17px;
    font-weight: 800;
}
.overview-highlight-label {
    font-size: 9px;
    font-weight: 800;
    letter-spacing: .10em;
    opacity: .86;
    margin-bottom: 3px;
}
.overview-highlight-title {
    font-size: 17px;
    font-weight: 800;
    line-height: 1.15;
}
.overview-highlight-text {
    font-size: 10.5px;
    line-height: 1.35;
    opacity: .86;
    margin-top: 3px;
}
@media (max-width: 900px) {
    .overview-highlights { grid-template-columns: 1fr; }
}

/* ========================= SECTION LABEL ========================= */
.section-title {
    color: #17324d;
    font-size: 20px;
    font-weight: 800;
    margin: 26px 0 12px;
}
.section-subtitle {
    color: #8193a2;
    font-size: 12px;
    margin: -6px 0 14px;
}

/* ========================= KPI CARDS ========================= */
div[data-testid="stMetric"] {
    background: rgba(255,255,255,.96);
    border: 1px solid #d9e8ee;
    border-radius: 18px;
    padding: 17px 19px;
    min-height: 112px;
    box-shadow: var(--shadow);
    transition: transform .18s ease, box-shadow .18s ease;
}
div[data-testid="stMetric"]:hover {
    transform: translateY(-2px);
    box-shadow: 0 18px 36px rgba(38,78,103,.13);
}
div[data-testid="stMetric"] label {
    color: #5f7587 !important;
    font-weight: 700 !important;
    font-size: 12px !important;
}
div[data-testid="stMetricValue"] {
    color: #17324d !important;
    font-weight: 800 !important;
    font-size: 27px !important;
}

/* ========================= CARDS / CHARTS / TABLES ========================= */
div[data-testid="stPlotlyChart"] {
    background: rgba(255,255,255,.96);
    border: 1px solid #d9e8ee;
    border-radius: 18px;
    padding: 7px 7px 0;
    box-shadow: var(--shadow);
    margin-bottom: 16px;
}
div[data-testid="stDataFrame"] {
    border: 1px solid #d9e8ee;
    border-radius: 16px;
    overflow: hidden;
    box-shadow: 0 8px 24px rgba(38,78,103,.07);
}
div[data-testid="stExpander"] {
    background: rgba(255,255,255,.94);
    border: 1px solid #d9e8ee;
    border-radius: 16px;
}
div[data-testid="stAlert"] {
    border-radius: 14px;
    border: 1px solid #d8e7ed;
}

/* insight cards */
.insight-card {
    background: rgba(255,255,255,.92);
    border: 1px solid #d9e8ee;
    border-left: 4px solid #159c9c;
    border-radius: 14px;
    padding: 12px 15px;
    margin: 7px 0;
    box-shadow: 0 8px 20px rgba(38,78,103,.055);
    color: #405b6d;
    font-size: 13px;
    line-height: 1.5;
}

/* status */
.status-card {
    background: rgba(255,255,255,.94);
    border: 1px solid #d9e8ee;
    border-radius: 18px;
    padding: 16px;
    box-shadow: 0 10px 28px rgba(38,78,103,.07);
}
.status-pill {
    display:inline-block;
    padding: 5px 9px;
    border-radius: 999px;
    font-size: 10px;
    font-weight: 800;
    letter-spacing:.04em;
}
.status-ok { background:#e2f7ef; color:#148463; }
.status-wait { background:#fff2d9; color:#b6760a; }

/* ========================= BUTTONS / UPLOAD ========================= */
.stButton > button,
.stDownloadButton > button {
    border-radius: 11px;
    border: 1px solid #cfe1e7;
    background: rgba(255,255,255,.96);
    color: #15566e;
    font-weight: 700;
    min-height: 41px;
    transition: .18s ease;
}
.stButton > button:hover,
.stDownloadButton > button:hover {
    border-color: #74bfc2;
    color: #087d82;
    box-shadow: 0 7px 18px rgba(21,156,156,.12);
}
[data-testid="stFileUploader"] {
    background: rgba(255,255,255,.78);
    border: 1.5px dashed #87bfc7;
    border-radius: 17px;
    padding: 10px;
}
button[data-baseweb="tab"] {
    font-weight: 700;
    color: #6e8291 !important;
}
button[data-baseweb="tab"][aria-selected="true"] {
    color: #0f8e91 !important;
}
hr { border-color: #dce8ed; }

/* ========================= EMPTY / INFO ========================= */
.data-note {
    display:flex;
    align-items:center;
    gap:10px;
    padding:10px 13px;
    border-radius:12px;
    background:rgba(255,255,255,.62);
    border:1px solid #dce8ed;
    color:#6d8190;
    font-size:12px;
}
</style>
""", unsafe_allow_html=True)


def page_hero(eyebrow, title, subtitle, show_decorative=False):
    """Hero section yang konsisten dengan desain dashboard.

    Decorative summary cards hanya ditampilkan pada Overview agar halaman
    Pelatihan, Bimbingan, dan Data Detail tetap fokus dan tidak repetitif.
    """
    st.markdown(
        f"""
        <div class="dashboard-hero">
            <div class="eyebrow">{eyebrow}</div>
            <h1>{title}</h1>
            <p>{subtitle}</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    if show_decorative:
        st.markdown(
            """
            <div class="overview-highlights">
                <div class="overview-highlight overview-highlight-gold">
                    <div class="overview-highlight-icon">◈</div>
                    <div>
                        <div class="overview-highlight-label">CAKUPAN PROGRAM</div>
                        <div class="overview-highlight-title">2 Area Monitoring</div>
                        <div class="overview-highlight-text">Pelatihan Produktivitas &amp; Bimbingan Konsultasi</div>
                    </div>
                </div>
                <div class="overview-highlight overview-highlight-teal">
                    <div class="overview-highlight-icon">✓</div>
                    <div>
                        <div class="overview-highlight-label">STATUS DATA</div>
                        <div class="overview-highlight-title">Dashboard Aktif</div>
                        <div class="overview-highlight-text">Data siap digunakan untuk monitoring dan analisis</div>
                    </div>
                </div>
                <div class="overview-highlight overview-highlight-blue">
                    <div class="overview-highlight-icon">2026</div>
                    <div>
                        <div class="overview-highlight-label">PERIODE</div>
                        <div class="overview-highlight-title">Tahun Berjalan</div>
                        <div class="overview-highlight-text">Ringkasan aktivitas berdasarkan data yang tersedia</div>
                    </div>
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )


def section_heading(title, subtitle=None):
    html = f'<div class="section-title">{title}</div>'
    if subtitle:
        html += f'<div class="section-subtitle">{subtitle}</div>'
    st.markdown(html, unsafe_allow_html=True)


def insight_card(text):
    st.markdown(
        f'<div class="insight-card">{text.replace("**", "")}</div>',
        unsafe_allow_html=True
    )


# =========================================================
# NAMA SHEET
# =========================================================

TRAIN_SHEET = "Pelatihan Produktivitas"
BIM_SHEET = "Bimbingan Konsultasi"
REKAP_SHEET = "Rekap Pel Prod"
REKAP_BIM_SHEET = "Rekap Bimkon"

MONTHS = [
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember"
]


# =========================================================
# BACA EXCEL
# =========================================================

def _canonical_header_name(value):
    """Samakan variasi nama kolom antar file Excel."""
    if pd.isna(value):
        return "Kolom"
    raw = str(value).strip()
    key = re.sub(r"[^a-z0-9]", "", raw.lower())

    aliases = {
        "namapeserta": "Nama Peserta",
        "statuskelulusan": "Status KeLulusan",
        "statuskelulusanpeserta": "Status KeLulusan",
        "statuslulus": "Status KeLulusan",
        "namalembaga": "Nama Lembaga",
        "judulprogrampelatihan": "Judul Program Pelatihan",
        "tanggalmulaipelatihan": "Tanggal Mulai Pelatihan",
        "tanggalselesaipelatihan": "Tanggal Selesai Pelatihan",
        "metodepelatihan": "Metode Pelatihan",
        "jenispelatihan": "Jenis Pelatihan",
        "statusselesaipelatihan": "Status Selesai Pelatihan",
        "kejuruan": "Kejuruan",
        "provinsi": "Provinsi",
        "kabkota": "Kab./Kota",
        "kabupatenkota": "Kab./Kota",
        "lembagainstansi": "Lembaga/Instansi",
        "targetpesertapelatihanp3": "Target Pelatihan P3",
        "targetperusahaan": "Target Perusahaan",
        "realisasi": "Realisasi",
        "namaperusahaan": "NAMA PERUSAHAAN",
        "bidangusaha": "BIDANG USAHA",
        "namakabupatenkota": "NAMA KABUPATEN/KOTA",
    }
    return aliases.get(key, raw)


def _coalesce_duplicate_columns(df):
    """Gabungkan kolom yang memiliki nama sama dengan mengambil nilai non-kosong pertama."""
    if df.empty:
        return df
    result = pd.DataFrame(index=df.index)
    for name in dict.fromkeys(df.columns):
        cols = [c for c in df.columns if c == name]
        if len(cols) == 1:
            result[name] = df[cols[0]]
        else:
            block = df.loc[:, cols]
            result[name] = block.bfill(axis=1).iloc[:, 0]
    result.columns = make_unique(result.columns)
    return result


def _prepare_multi_file_sheet(raw, sheet):
    """Menyiapkan sheet per file berdasarkan nama header, bukan posisi kolom."""
    if raw is None or raw.empty:
        return pd.DataFrame()

    header = None
    for i, row in raw.iterrows():
        vals = [str(v).strip().upper() for v in row.tolist() if pd.notna(v)]
        joined = " | ".join(vals)
        if sheet == TRAIN_SHEET:
            if "NAMA PESERTA" in vals and any("STATUS" in v and "LULUS" in v for v in vals):
                header = i
                break
        elif sheet == BIM_SHEET:
            if "NAMA PERUSAHAAN" in vals:
                header = i
                break
        elif sheet in [REKAP_SHEET, REKAP_BIM_SHEET]:
            if "LEMBAGA/INSTANSI" in joined and "REALISASI" in joined:
                header = i
                break

    if header is None:
        return raw.copy()

    columns = [_canonical_header_name(v) for v in raw.iloc[header].tolist()]
    data = raw.iloc[header + 1:].copy()
    data.columns = columns
    data = data.dropna(how="all")
    data = _coalesce_duplicate_columns(data)

    if sheet == TRAIN_SHEET and "Nama Peserta" in data.columns:
        data = data[data["Nama Peserta"].astype(str).str.strip().str.upper() != "NAMA PESERTA"]
    if sheet == BIM_SHEET and "NAMA PERUSAHAAN" in data.columns:
        data = data[data["NAMA PERUSAHAAN"].astype(str).str.strip().str.upper() != "NAMA PERUSAHAAN"]

    # Header + data agar fungsi cleaning yang sudah ada tetap kompatibel.
    header_row = pd.DataFrame([list(data.columns)], columns=data.columns)
    return pd.concat([header_row, data], ignore_index=True)


@st.cache_data(show_spinner=False)
def load_excel(file_items):
    """Membaca beberapa file Excel dan menggabungkan data berdasarkan nama kolom."""
    sheets = {}

    for file_name, file_bytes in file_items:
        xls = pd.ExcelFile(BytesIO(file_bytes), engine="openpyxl")

        for sheet in xls.sheet_names:
            raw = pd.read_excel(xls, sheet_name=sheet, header=None)
            prepared = _prepare_multi_file_sheet(raw, sheet)

            if sheet not in sheets or sheets[sheet].empty:
                sheets[sheet] = prepared
                continue

            # Gabungkan berdasarkan nama kolom yang sudah dinormalisasi.
            base = sheets[sheet]
            if sheet in [TRAIN_SHEET, BIM_SHEET]:
                base_header = [_canonical_header_name(v) for v in base.iloc[0].tolist()]
                base_data = base.iloc[1:].copy()
                base_data.columns = base_header
                base_data = _coalesce_duplicate_columns(base_data)

                new_header = [_canonical_header_name(v) for v in prepared.iloc[0].tolist()]
                new_data = prepared.iloc[1:].copy()
                new_data.columns = new_header
                new_data = _coalesce_duplicate_columns(new_data)

                all_columns = list(dict.fromkeys(list(base_data.columns) + list(new_data.columns)))
                base_data = base_data.reindex(columns=all_columns)
                new_data = new_data.reindex(columns=all_columns)
                combined = pd.concat([base_data, new_data], ignore_index=True, sort=False)
                sheets[sheet] = pd.concat(
                    [pd.DataFrame([list(combined.columns)], columns=combined.columns), combined],
                    ignore_index=True
                )
            else:
                # Rekap tetap mengikuti struktur sumber, tetapi cegah duplicate columns.
                base_header = [_canonical_header_name(v) for v in base.iloc[0].tolist()]
                base_data = base.iloc[1:].copy()
                base_data.columns = base_header
                new_header = [_canonical_header_name(v) for v in prepared.iloc[0].tolist()]
                new_data = prepared.iloc[1:].copy()
                new_data.columns = new_header
                all_columns = list(dict.fromkeys(list(base_data.columns) + list(new_data.columns)))
                base_data = base_data.reindex(columns=all_columns)
                new_data = new_data.reindex(columns=all_columns)
                combined = pd.concat([base_data, new_data], ignore_index=True, sort=False)
                sheets[sheet] = pd.concat(
                    [pd.DataFrame([list(combined.columns)], columns=combined.columns), combined],
                    ignore_index=True
                )

    return sheets


# =========================================================
# KOLOM UNIK
# =========================================================

def make_unique(cols):
    """Buat nama kolom yang benar-benar unik, termasuk jika sumber
    sudah memiliki nama seperti A, A_2, A, atau kolom kosong.
    """
    used = set()
    counters = {}
    output = []

    for col in cols:
        if pd.isna(col):
            base = "Kolom"
        else:
            base = str(col).strip()
            if not base or base.lower() in {"nan", "none"}:
                base = "Kolom"

        # Pertahankan nama asli bila belum dipakai.
        if base not in used:
            candidate = base
            counters.setdefault(base, 1)
        else:
            # Cari suffix berikutnya yang juga belum dipakai.
            n = counters.get(base, 1) + 1
            candidate = f"{base}_{n}"
            while candidate in used:
                n += 1
                candidate = f"{base}_{n}"
            counters[base] = n

        used.add(candidate)
        output.append(candidate)

    return output


# =========================================================
# NORMALISASI TEKS
# =========================================================

def normalisasi_teks(value):

    if pd.isna(value):
        return pd.NA

    text = str(value).strip()

    if text == "":
        return pd.NA

    return text


# =========================================================
# STANDARDISASI PROVINSI
# =========================================================

def standardisasi_provinsi(value):

    if pd.isna(value):
        return "Tidak Diisi"

    text = str(value).strip().upper()

    mapping = {

        "DKI JAKARTA": "DKI Jakarta",

        "JAWA BARAT": "Jawa Barat",
        "JAWA TENGAH": "Jawa Tengah",
        "JAWA TIMUR": "Jawa Timur",

        "BANTEN": "Banten",

        "SUMATERA BARAT": "Sumatera Barat",
        "SUMATERA UTARA": "Sumatera Utara",
        "SUMATERA SELATAN": "Sumatera Selatan",

        "KALIMANTAN TIMUR": "Kalimantan Timur",
        "KALIMANTAN BARAT": "Kalimantan Barat",
        "KALIMANTAN SELATAN": "Kalimantan Selatan",
        "KALIMANTAN TENGAH": "Kalimantan Tengah",
        "KALIMANTAN UTARA": "Kalimantan Utara",

        "ACEH": "Aceh",

        "SULAWESI SELATAN": "Sulawesi Selatan",
        "SULSEL": "Sulawesi Selatan",

        "SULAWESI TENGGARA": "Sulawesi Tenggara",
        "SULAWESI UTARA": "Sulawesi Utara",
        "SULAWESI TENGAH": "Sulawesi Tengah",

        "PAPUA BARAT DAYA": "Papua Barat Daya",
        "PAPUA": "Papua"
    }

    return mapping.get(
        text,
        str(value).strip().title()
    )


# =========================================================
# STANDARDISASI METODE
# =========================================================

def standardisasi_metode(value):

    if pd.isna(value):
        return "Tidak Diisi"

    text = str(value).strip().lower()

    if text == "online":
        return "Online"

    if text == "offline":
        return "Offline"

    if text == "hybrid":
        return "Hybrid"

    return str(value).strip().title()


# =========================================================
# JENIS PELATIHAN
# =========================================================

def standardisasi_jenis_pelatihan(value):

    if pd.isna(value):
        return "Tidak Diisi"

    text = str(value).strip().upper()

    if text == "NON BOARDING":
        return "Non Boarding"

    if text == "BOARDING":
        return "Boarding"

    return str(value).strip().title()


# =========================================================
# STATUS LULUS
# =========================================================

def standardisasi_status_lulus(value):

    if pd.isna(value):
        return "Tidak Diisi"

    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)

    # Variasi penulisan yang umum muncul antar file Excel.
    if text in {"lulus", "lulus.", "100", "1", "ya", "yes", "passed", "kompeten"}:
        return "Lulus"

    if text in {"tidak lulus", "tidak lulus.", "0", "tidak", "no", "not passed", "tidak kompeten"}:
        return "Tidak Lulus"

    if "tidak lulus" in text or "tidak kompeten" in text:
        return "Tidak Lulus"
    if text.startswith("lulus") or "kompeten" == text:
        return "Lulus"

    return str(value).strip().title()


# =========================================================
# STATUS SELESAI
# =========================================================

def standardisasi_status_selesai(value):

    if pd.isna(value):
        return "Tidak Diisi"

    text = str(value).strip().lower()

    if text in [
        "selesai",
        "ya",
        "1"
    ]:
        return "Selesai"

    if text in [
        "belum selesai",
        "tidak",
        "0"
    ]:
        return "Belum Selesai"

    return str(value).strip().title()


# =========================================================
# KEJURUAN
# =========================================================

def standardisasi_kejuruan(value):

    if pd.isna(value):
        return "Tidak Diisi"

    text = str(value).strip().lower()

    if text == "produktivitas":
        return "Produktivitas"

    return str(value).strip().title()


# =========================================================
# KATEGORI BIDANG USAHA
# =========================================================

def kategori_bidang_usaha(value):

    if pd.isna(value):
        return "Tidak Diisi"

    text = str(value).strip().lower()

    if text == "":
        return "Tidak Diisi"


    # MAKANAN & MINUMAN

    if any(
        k in text
        for k in [

            "kuliner",
            "bakery",
            "keripik",
            "snack",
            "dimsum",
            "kue",
            "catering",
            "bawang goreng",
            "jus",
            "bubur",
            "jamu",
            "peyek",
            "tape",
            "bandeng presto",
            "sambal",
            "seblak",
            "warung makan",
            "minuman",
            "ayam gepuk",
            "donat",
            "abon",
            "crispy",
            "tepung",
            "makanan"
        ]
    ):

        return "Makanan & Minuman"


    # TEKSTIL

    if any(
        k in text
        for k in [

            "menjahit",
            "jahit",
            "fashion",
            "konveksi",
            "pakaian",
            "tekstil",
            "bordir"
        ]
    ):

        return "Tekstil & Fashion"


    # OTOMOTIF

    if any(
        k in text
        for k in [

            "bengkel",
            "otomotif",
            "motor",
            "mobil",
            "kendaraan"
        ]
    ):

        return "Otomotif"


    # PERTANIAN & PERIKANAN

    if any(
        k in text
        for k in [

            "budidaya",
            "pertanian",
            "peternakan",
            "perikanan",
            "ikan",
            "perkebunan"
        ]
    ):

        return "Pertanian & Perikanan"


    # PERDAGANGAN

    if any(
        k in text
        for k in [

            "dagang",
            "toko",
            "agen",
            "perdagangan",
            "distributor"
        ]
    ):

        return "Perdagangan"


    # JASA

    if any(
        k in text
        for k in [

            "jasa",
            "expedisi",
            "ekspedisi",
            "service"
        ]
    ):

        return "Jasa"


    return "Lainnya"


# =========================================================
# CLEANING REKAP REALISASI LEMBAGA
# =========================================================

def clean_rekap_realisasi(raw):
    df = raw.copy()
    header = None

    for i, row in df.iterrows():
        values = [
            str(v).strip().upper()
            for v in row.tolist()
            if pd.notna(v)
        ]
        joined = " | ".join(values)
        if "LEMBAGA/INSTANSI" in joined and "TARGET PESERTA PELATIHAN P3" in joined:
            header = i
            break

    if header is None:
        return pd.DataFrame()

    df.columns = make_unique(df.iloc[header])
    df = df.iloc[header + 1:].copy().dropna(how="all")

    rename_map = {}
    realisasi_found = False
    for col in df.columns:
        text = str(col).strip().upper()
        if "LEMBAGA/INSTANSI" in text:
            rename_map[col] = "Lembaga/Instansi"
        elif "TARGET PESERTA PELATIHAN P3" in text and "PEKERJAAN HIJAU" not in text:
            rename_map[col] = "Target Pelatihan P3"
        elif text == "REALISASI" and not realisasi_found:
            rename_map[col] = "Realisasi Pelatihan P3"
            realisasi_found = True

    df = df.rename(columns=rename_map)
    if "Lembaga/Instansi" not in df.columns or "Realisasi Pelatihan P3" not in df.columns:
        return pd.DataFrame()

    df["Lembaga/Instansi"] = df["Lembaga/Instansi"].apply(normalisasi_teks)
    df = df[df["Lembaga/Instansi"].notna()].copy()
    df = df[~df["Lembaga/Instansi"].astype(str).str.upper().isin([
        "PELATIHAN PRODUKTIVITAS", "BIMBINGAN KONSULTANSI", "JUMLAH TOTAL", "TOTAL"
    ])]

    for col in ["Target Pelatihan P3", "Realisasi Pelatihan P3"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

    df["Status Realisasi"] = df["Realisasi Pelatihan P3"].apply(
        lambda x: "Sudah Terealisasi" if x > 0 else "Belum Terealisasi"
    )

    if "Target Pelatihan P3" in df.columns:
        df["Persentase Realisasi"] = df.apply(
            lambda r: (r["Realisasi Pelatihan P3"] / r["Target Pelatihan P3"] * 100)
            if r["Target Pelatihan P3"] > 0 else 0,
            axis=1
        ).round(1)

    return df.reset_index(drop=True)


# =========================================================
# CLEANING REKAP REALISASI BIMBINGAN KONSULTASI
# =========================================================

def clean_rekap_bimkon(raw):
    df = raw.copy()
    header = None

    # Cari baris header yang berisi kolom lembaga, target, dan realisasi.
    for i, row in df.iterrows():
        values = [
            str(v).strip().upper()
            for v in row.tolist()
            if pd.notna(v)
        ]
        joined = " | ".join(values)

        if (
            "LEMBAGA/INSTANSI" in joined
            and "TARGET PERUSAHAAN" in joined
            and "REALISASI" in joined
        ):
            header = i
            break

    if header is None:
        return pd.DataFrame()

    header_values = [
        str(v).strip() if pd.notna(v) else "Kolom"
        for v in raw.iloc[header].tolist()
    ]
    df = raw.iloc[header + 1:].copy()
    df.columns = make_unique(header_values)
    df = df.dropna(how="all")

    rename_map = {}
    for col in df.columns:
        text = str(col).strip().upper()
        if text == "LEMBAGA/INSTANSI":
            rename_map[col] = "Lembaga/Instansi"
        elif text == "TARGET PERUSAHAAN":
            rename_map[col] = "Target Perusahaan"
        elif text == "REALISASI":
            rename_map[col] = "Realisasi Bimbingan"

    df = df.rename(columns=rename_map)

    required = [
        "Lembaga/Instansi",
        "Target Perusahaan",
        "Realisasi Bimbingan"
    ]
    if not all(col in df.columns for col in required):
        return pd.DataFrame()

    df["Lembaga/Instansi"] = df["Lembaga/Instansi"].apply(normalisasi_teks)
    df = df[df["Lembaga/Instansi"].notna()].copy()

    exclude = [
        "BIMBINGAN KONSULTANSI",
        "JUMLAH TOTAL",
        "TOTAL",
        "PELATIHAN PRODUKTIVITAS"
    ]
    df = df[
        ~df["Lembaga/Instansi"].astype(str).str.strip().str.upper().isin(exclude)
    ].copy()

    for col in ["Target Perusahaan", "Realisasi Bimbingan"]:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

    df["Status Realisasi"] = df["Realisasi Bimbingan"].apply(
        lambda x: "Sudah Terealisasi" if x > 0 else "Belum Terealisasi"
    )

    df["Persentase Realisasi"] = df.apply(
        lambda r: (
            r["Realisasi Bimbingan"] / r["Target Perusahaan"] * 100
        ) if r["Target Perusahaan"] > 0 else 0,
        axis=1
    ).round(1)

    return df.reset_index(drop=True)


# =========================================================
# CLEANING PELATIHAN
# =========================================================

def clean_training(raw):

    df = raw.copy()

    # Normalisasi nama header agar multi-file tidak membuat Status KeLulusan
    # terpecah menjadi kolom berbeda karena variasi kapitalisasi/ejaan.
    normalized_headers = [
        _canonical_header_name(v)
        for v in df.iloc[0].tolist()
    ]
    df = df.iloc[1:].copy()
    df.columns = normalized_headers
    df = _coalesce_duplicate_columns(df)

    df = df.dropna(
        how="all"
    )


    # Hanya peserta

    if "Nama Peserta" in df.columns:

        df = df[
            df["Nama Peserta"].notna()
        ].copy()


    # Tanggal

    for col in [

        "Tanggal Mulai Pelatihan",
        "Tanggal Selesai Pelatihan",
        "Tanggal Lahir"

    ]:

        if col in df.columns:

            df[col] = pd.to_datetime(
                df[col],
                errors="coerce",
                dayfirst=True
            )


    # Numeric

    for col in [

        "Durasi (JP)",
        "Absensi"

    ]:

        if col in df.columns:

            df[col] = pd.to_numeric(
                df[col],
                errors="coerce"
            )


    # Bulan & Tahun

    if "Tanggal Mulai Pelatihan" in df.columns:

        bulan = {

            1: "Januari",
            2: "Februari",
            3: "Maret",
            4: "April",
            5: "Mei",
            6: "Juni",
            7: "Juli",
            8: "Agustus",
            9: "September",
            10: "Oktober",
            11: "November",
            12: "Desember"
        }

        df["Bulan Pelatihan"] = (
            df["Tanggal Mulai Pelatihan"]
            .dt.month
            .map(bulan)
        )

        df["Tahun Pelatihan"] = (
            df["Tanggal Mulai Pelatihan"]
            .dt.year
        )


    # Standardisasi

    if "Provinsi" in df.columns:

        df["Provinsi"] = (
            df["Provinsi"]
            .apply(standardisasi_provinsi)
        )


    if "Metode Pelatihan" in df.columns:

        df["Metode Pelatihan"] = (
            df["Metode Pelatihan"]
            .apply(standardisasi_metode)
        )


    if "Jenis Pelatihan" in df.columns:

        df["Jenis Pelatihan"] = (
            df["Jenis Pelatihan"]
            .apply(standardisasi_jenis_pelatihan)
        )


    if "Status KeLulusan" in df.columns:

        df["Status KeLulusan"] = (
            df["Status KeLulusan"]
            .apply(standardisasi_status_lulus)
        )


    if "Status Selesai Pelatihan" in df.columns:

        df["Status Selesai Pelatihan"] = (
            df["Status Selesai Pelatihan"]
            .apply(standardisasi_status_selesai)
        )


    if "Kejuruan" in df.columns:

        df["Kejuruan"] = (
            df["Kejuruan"]
            .apply(standardisasi_kejuruan)
        )


    # Normalisasi teks

    for col in [

        "Kab./Kota",
        "Nama Lembaga",
        "Judul Program Pelatihan",
        "Jenis Kelamin",
        "Pendidikan Terakhir (setara dengan)"

    ]:

        if col in df.columns:

            df[col] = (
                df[col]
                .apply(normalisasi_teks)
            )


    return df.reset_index(
        drop=True
    )


# =========================================================
# CLEANING BIMBINGAN
# =========================================================

def clean_bim(raw):

    df = raw.copy()

    header = 0


    # Cari header

    for i, row in df.iterrows():

        values = (
            row.astype(str)
            .str.upper()
            .tolist()
        )

        if "NAMA PERUSAHAAN" in values:

            header = i

            break


    df.columns = make_unique(
        df.iloc[header]
    )

    df = df.iloc[
        header + 1:
    ].copy()

    df = df.dropna(
        how="all"
    )


    # Hanya perusahaan

    if "NAMA PERUSAHAAN" in df.columns:

        df = df[
            df["NAMA PERUSAHAAN"].notna()
        ].copy()

        df = df[
            df["NAMA PERUSAHAAN"]
            .astype(str)
            .str.upper()
            != "CONTOH PENGISIAN"
        ]


    # Tanggal pelaksanaan

    if "WAKTU PELAKSANAAN" in df.columns:

        df["Tanggal Pelaksanaan"] = pd.to_datetime(
            df["WAKTU PELAKSANAAN"],
            errors="coerce",
            dayfirst=True
        )


        # Contoh:
        # 20/7/2026 s/d 12/08/2026

        mask = (
            df["Tanggal Pelaksanaan"]
            .isna()
        )

        if mask.any():

            extracted = (
                df.loc[
                    mask,
                    "WAKTU PELAKSANAAN"
                ]
                .astype(str)
                .str.extract(
                    r"(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})"
                )[0]
            )

            df.loc[
                mask,
                "Tanggal Pelaksanaan"
            ] = pd.to_datetime(
                extracted,
                errors="coerce",
                dayfirst=True
            )


    # Tanggal data

    if (
        "DATA PER TANGGAL (tanggal/bulan/tahun)"
        in df.columns
    ):

        df["Tanggal Data"] = pd.to_datetime(
            df[
                "DATA PER TANGGAL (tanggal/bulan/tahun)"
            ],
            errors="coerce",
            dayfirst=True
        )


    # Bulan & tahun

    if "Tanggal Pelaksanaan" in df.columns:

        bulan = {

            1: "Januari",
            2: "Februari",
            3: "Maret",
            4: "April",
            5: "Mei",
            6: "Juni",
            7: "Juli",
            8: "Agustus",
            9: "September",
            10: "Oktober",
            11: "November",
            12: "Desember"
        }

        df["Bulan Bimbingan"] = (
            df["Tanggal Pelaksanaan"]
            .dt.month
            .map(bulan)
        )

        df["Tahun Bimbingan"] = (
            df["Tanggal Pelaksanaan"]
            .dt.year
        )


    # Normalisasi teks

    for col in [

        "NAMA KABUPATEN/KOTA",
        "BIDANG USAHA",
        "PELAKSANA",
        "AKTIVITAS PENINGKATAN PRODUKTIVITAS",
        "HASIL PENINGKATAN PRODUKTIVITAS"

    ]:

        if col in df.columns:

            df[col] = (
                df[col]
                .apply(normalisasi_teks)
            )


    # Bidang usaha

    if "BIDANG USAHA" in df.columns:

        df["BIDANG USAHA ASLI"] = (
            df["BIDANG USAHA"]
        )

        df["Bidang Usaha Kategori"] = (
            df["BIDANG USAHA"]
            .apply(kategori_bidang_usaha)
        )


    return df.reset_index(
        drop=True
    )


# =========================================================
# FILTER
# =========================================================

def selectbox_filter(
    df,
    col,
    label,
    key=None
):

    if col not in df.columns:
        return df

    values = (
        df[col]
        .dropna()
        .astype(str)
        .str.strip()
    )
    values = sorted(values[values != ""].unique().tolist())

    selected = st.sidebar.selectbox(
        label,
        ["Semua"] + values,
        index=0,
        key=key
    )

    if selected == "Semua":
        return df

    return df[
        df[col].astype(str).str.strip() == selected
    ].copy()


# =========================================================
# FORMAT ANGKA
# =========================================================

def format_number(value):

    try:

        return f"{int(value):,}".replace(
            ",",
            "."
        )

    except:

        return "0"


# =========================================================
# PERSENTASE
# =========================================================

def percentage(part, total):

    if total == 0:
        return 0

    return (
        part / total
    ) * 100


# =========================================================
# DATA BULANAN
# =========================================================

def monthly_data(
    df,
    month_column,
    value_name
):

    if month_column not in df.columns:

        return pd.DataFrame({
            "Bulan": MONTHS,
            value_name: [0] * 12
        })


    x = (
        df[month_column]
        .value_counts()
        .reindex(MONTHS)
        .fillna(0)
        .reset_index()
    )

    x.columns = [
        "Bulan",
        value_name
    ]

    return x


# =========================================================
# STYLE PLOTLY
# =========================================================

def style_chart(
    fig,
    height=330,
    horizontal=False,
    show_legend=False
):
    """Style Plotly yang konsisten: clean, lega, dan aman untuk label."""
    fig.update_layout(
        height=height,
        margin=dict(
            l=120 if horizontal else 18,
            r=52,
            t=54,
            b=58
        ),
        paper_bgcolor="rgba(255,255,255,0)",
        plot_bgcolor="rgba(255,255,255,0)",
        font=dict(family="Arial", size=11, color="#17324d"),
        title=dict(
            x=0.02, xanchor="left",
            font=dict(size=14, color="#17324d")
        ),
        legend=dict(
            orientation="h",
            yanchor="bottom", y=1.01,
            x=0.02, xanchor="left",
            font=dict(size=10),
            bgcolor="rgba(255,255,255,0)"
        ),
        hoverlabel=dict(
            bgcolor="white", font_size=12, font_family="Arial",
            bordercolor="#d9e8ee"
        ),
        bargap=0.34,
        uniformtext_minsize=9,
        uniformtext_mode="hide"
    )

    fig.update_xaxes(
        showgrid=True, gridcolor="#e8eef4", zeroline=False, showline=False,
        tickfont=dict(size=10, color="#71859a"),
        title_font=dict(size=10, color="#71859a"),
        automargin=True
    )
    fig.update_yaxes(
        showgrid=True, gridcolor="#e8eef4", zeroline=False, showline=False,
        tickfont=dict(size=10, color="#71859a"),
        title_font=dict(size=10, color="#71859a"),
        automargin=True
    )

    if not show_legend:
        fig.update_layout(showlegend=False)

    return fig

def clean_chart_text(values):
    """Tampilkan label hanya untuk nilai yang terisi agar grafik tidak penuh angka 0."""
    return [
        format_number(v) if pd.notna(v) and float(v) > 0 else ""
        for v in values
    ]


def short_month_labels():
    return [
        "Jan", "Feb", "Mar", "Apr", "Mei", "Jun",
        "Jul", "Agu", "Sep", "Okt", "Nov", "Des"
    ]


def show_kpi_tags(values, title, icon="📌"):
    """Tampilkan rincian KPI sebagai chip/tag yang rapi dan responsif."""
    cleaned = sorted({str(v).strip() for v in values if pd.notna(v) and str(v).strip()})
    if not cleaned:
        return

    tags = "".join(
        '<span style="display:inline-block; flex:1 1 30%; min-width:220px; '
        'padding:9px 12px; border:1px solid #dce8f1; border-radius:10px; '
        'background:#f7fbfe; color:#35546d; font-size:11px; line-height:1.35; '
        'box-sizing:border-box;">' + html.escape(v) + '</span>'
        for v in cleaned
    )

    st.markdown(
        '<div style="margin:4px 0 18px 0; padding:14px 16px 16px 16px; '
        'border:1px solid #dce8f1; border-radius:14px; background:rgba(255,255,255,.72);">'
        '<div style="font-size:12px; font-weight:700; color:#17324d; margin-bottom:10px;">'
        + icon + ' ' + html.escape(title) + ' <span style="font-weight:500; color:#7a8ea0;">('
        + str(len(cleaned)) + ')</span></div>'
        '<div style="display:flex; flex-wrap:wrap; gap:8px; align-items:stretch;">'
        + tags + '</div></div>',
        unsafe_allow_html=True
    )


# =========================================================
# INSIGHT PELATIHAN
# =========================================================

def insights_training(df):

    insights = []

    if df.empty:

        return [
            "Tidak ada data sesuai filter."
        ]


    total = len(df)


    # Provinsi

    if "Provinsi" in df.columns:

        x = (
            df["Provinsi"]
            .fillna("Tidak Diisi")
            .value_counts()
        )

        if not x.empty:

            nama = x.index[0]
            jumlah = int(x.iloc[0])
            persen = percentage(
                jumlah,
                total
            )

            insights.append(
                (
                    "🏆 Provinsi Dominan — "
                    f"{nama} memiliki peserta terbanyak, "
                    f"yaitu {format_number(jumlah)} peserta "
                    f"({persen:.1f}%)."
                )
            )


    # Bulan

    if "Bulan Pelatihan" in df.columns:

        x = (
            df["Bulan Pelatihan"]
            .value_counts()
            .reindex(MONTHS)
            .fillna(0)
        )

        if x.max() > 0:

            bulan = x.idxmax()
            jumlah = int(x.max())

            insights.append(
                (
                    "📈 Bulan Tertinggi — "
                    f"Jumlah peserta tertinggi terjadi pada "
                    f"{bulan}, sebanyak "
                    f"{format_number(jumlah)} peserta."
                )
            )


    # Kelulusan

    if "Status KeLulusan" in df.columns:

        lulus = (
            df["Status KeLulusan"]
            .astype(str)
            .str.lower()
            == "lulus"
        ).sum()

        persen = percentage(
            lulus,
            total
        )

        insights.append(
            (
                "🎓 Kelulusan — "
                f"Sebanyak {format_number(lulus)} peserta "
                f"berstatus lulus dengan tingkat kelulusan "
                f"{persen:.1f}%."
            )
        )


    # Metode

    if "Metode Pelatihan" in df.columns:

        x = (
            df["Metode Pelatihan"]
            .fillna("Tidak Diisi")
            .value_counts()
        )

        if not x.empty:

            metode = x.index[0]
            jumlah = int(x.iloc[0])
            persen = percentage(
                jumlah,
                total
            )

            insights.append(
                (
                    "⚙️ Metode Dominan — "
                    f"Metode {metode} digunakan oleh "
                    f"{format_number(jumlah)} peserta "
                    f"({persen:.1f}%)."
                )
            )


    return insights


# =========================================================
# INSIGHT BIMBINGAN
# =========================================================

def insights_bim(df):

    insights = []

    if df.empty:

        return [
            "Tidak ada data sesuai filter."
        ]


    total = len(df)


    # Bidang usaha

    if "Bidang Usaha Kategori" in df.columns:

        x = (
            df["Bidang Usaha Kategori"]
            .fillna("Tidak Diisi")
            .value_counts()
        )

        if not x.empty:

            bidang = x.index[0]
            jumlah = int(x.iloc[0])
            persen = percentage(
                jumlah,
                total
            )

            insights.append(
                (
                    "🏭 Bidang Usaha Dominan — "
                    f"{bidang} menjadi kategori dengan "
                    f"bimbingan terbanyak, yaitu "
                    f"{format_number(jumlah)} kegiatan "
                    f"({persen:.1f}%)."
                )
            )


    # Wilayah

    if "NAMA KABUPATEN/KOTA" in df.columns:

        x = (
            df["NAMA KABUPATEN/KOTA"]
            .fillna("Tidak Diisi")
            .value_counts()
        )

        if not x.empty:

            wilayah = x.index[0]
            jumlah = int(x.iloc[0])

            insights.append(
                (
                    "📍 Wilayah Dominan — "
                    f"{wilayah} menjadi wilayah dengan "
                    f"kegiatan bimbingan terbanyak, "
                    f"yaitu {format_number(jumlah)} kegiatan."
                )
            )


    # Bulan

    if "Bulan Bimbingan" in df.columns:

        x = (
            df["Bulan Bimbingan"]
            .value_counts()
            .reindex(MONTHS)
            .fillna(0)
        )

        if x.max() > 0:

            bulan = x.idxmax()
            jumlah = int(x.max())

            insights.append(
                (
                    "📅 Bulan Tertinggi — "
                    f"Aktivitas bimbingan tertinggi terjadi "
                    f"pada {bulan}, sebanyak "
                    f"{format_number(jumlah)} kegiatan."
                )
            )


    return insights


# =========================================================
# EXPORT PDF HASIL DASHBOARD
# =========================================================

def _pdf_text(text):
    """Bersihkan markdown sederhana dari insight sebelum dimasukkan ke PDF."""
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", str(text))
    return text.replace("—", "-").replace("📌", "").strip()


def _save_chart_png(fig):
    """Simpan matplotlib figure ke memory agar bisa dimasukkan ke PDF."""
    buf = BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _matplotlib_training_charts(df):
    charts = []
    # Tren peserta
    x = monthly_data(df, "Bulan Pelatihan", "Peserta")
    if not x.empty:
        fig, ax = plt.subplots(figsize=(8.0, 3.8))
        ax.plot(x["Bulan"], x["Peserta"], marker="o")
        ax.set_title("Peserta per Bulan")
        ax.set_xlabel("Bulan")
        ax.set_ylabel("Peserta")
        ax.tick_params(axis="x", rotation=45)
        charts.append(_save_chart_png(fig))

    # Top provinsi
    if "Provinsi" in df.columns:
        x = df["Provinsi"].fillna("Tidak Diisi").value_counts().head(10).sort_values()
        if not x.empty:
            fig, ax = plt.subplots(figsize=(8.0, 4.2))
            ax.barh(x.index.astype(str), x.values)
            ax.set_title("Top 10 Provinsi")
            ax.set_xlabel("Peserta")
            charts.append(_save_chart_png(fig))

    # Kelulusan
    if "Status KeLulusan" in df.columns:
        x = df["Status KeLulusan"].fillna("Tidak Diisi").value_counts()
        if not x.empty:
            fig, ax = plt.subplots(figsize=(6.5, 4.0))
            ax.pie(x.values, labels=x.index.astype(str), autopct="%1.1f%%")
            ax.set_title("Status Kelulusan")
            charts.append(_save_chart_png(fig))

    # Metode
    if "Metode Pelatihan" in df.columns:
        x = df["Metode Pelatihan"].fillna("Tidak Diisi").value_counts()
        if not x.empty:
            fig, ax = plt.subplots(figsize=(7.0, 4.0))
            ax.bar(x.index.astype(str), x.values)
            ax.set_title("Metode Pelatihan")
            ax.set_ylabel("Peserta")
            ax.tick_params(axis="x", rotation=25)
            charts.append(_save_chart_png(fig))
    return charts


def _matplotlib_bim_charts(df):
    charts = []
    # Tren bimbingan
    x = monthly_data(df, "Bulan Bimbingan", "Kegiatan")
    if not x.empty:
        fig, ax = plt.subplots(figsize=(8.0, 3.8))
        ax.plot(x["Bulan"], x["Kegiatan"], marker="o")
        ax.set_title("Tren Bimbingan Konsultasi")
        ax.set_xlabel("Bulan")
        ax.set_ylabel("Kegiatan")
        ax.tick_params(axis="x", rotation=45)
        charts.append(_save_chart_png(fig))

    # Bidang usaha
    if "Bidang Usaha Kategori" in df.columns:
        x = df["Bidang Usaha Kategori"].fillna("Tidak Diisi").value_counts().sort_values()
        if not x.empty:
            fig, ax = plt.subplots(figsize=(8.0, 4.2))
            ax.barh(x.index.astype(str), x.values)
            ax.set_title("Kegiatan Berdasarkan Bidang Usaha")
            ax.set_xlabel("Kegiatan")
            charts.append(_save_chart_png(fig))

    # Wilayah
    for col in ["NAMA KABUPATEN/KOTA", "Kabupaten/Kota", "NAMA KABUPATEN / KOTA"]:
        if col in df.columns:
            x = df[col].fillna("Tidak Diisi").value_counts().head(10).sort_values()
            if not x.empty:
                fig, ax = plt.subplots(figsize=(8.0, 4.2))
                ax.barh(x.index.astype(str), x.values)
                ax.set_title("Top 10 Wilayah")
                ax.set_xlabel("Kegiatan")
                charts.append(_save_chart_png(fig))
            break
    return charts


def create_dashboard_pdf(training_df=None, bim_df=None, rekap_df=None, rekap_bim_df=None, title="Hasil Analisis Dashboard Produktivitas 2026"):
    """Membuat PDF visual dari hasil dashboard sesuai filter aktif."""
    output = BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=A4,
        rightMargin=1.4*cm, leftMargin=1.4*cm,
        topMargin=1.3*cm, bottomMargin=1.3*cm
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PdfTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=17, leading=21, spaceAfter=8))
    styles.add(ParagraphStyle(name="PdfSub", parent=styles["Normal"], alignment=TA_CENTER, fontSize=9, textColor=colors.grey, spaceAfter=12))
    styles.add(ParagraphStyle(name="PdfHead", parent=styles["Heading2"], fontSize=13, leading=16, spaceBefore=8, spaceAfter=7))
    styles.add(ParagraphStyle(name="PdfBody", parent=styles["BodyText"], fontSize=9, leading=13, spaceAfter=5))
    styles.add(ParagraphStyle(name="PdfSmall", parent=styles["BodyText"], fontSize=8, leading=11, spaceAfter=4))

    story = [Paragraph(title, styles["PdfTitle"]),
             Paragraph("Laporan visual berdasarkan data dan filter yang sedang aktif pada dashboard.", styles["PdfSub"])]

    def add_kpis(items):
        data = [[Paragraph(f"<b>{k}</b><br/>{v}", styles["PdfSmall"]) for k, v in items]]
        table = Table(data, colWidths=[(A4[0]-2.8*cm)/len(items)]*len(items))
        table.setStyle(TableStyle([
            ("GRID", (0,0), (-1,-1), 0.5, colors.lightgrey),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("ALIGN", (0,0), (-1,-1), "CENTER"),
            ("TOPPADDING", (0,0), (-1,-1), 8),
            ("BOTTOMPADDING", (0,0), (-1,-1), 8),
        ]))
        story.extend([table, Spacer(1, 8)])

    if training_df is not None and not training_df.empty:
        df = training_df
        story.append(Paragraph("🎓 Pelatihan Produktivitas", styles["PdfHead"]))
        add_kpis([
            ("Total Peserta", format_number(len(df))),
            ("Total Program", format_number(df["Judul Program Pelatihan"].nunique() if "Judul Program Pelatihan" in df.columns else 0)),
            ("Total Provinsi", format_number(df["Provinsi"].nunique() if "Provinsi" in df.columns else 0)),
            ("Total Lembaga", format_number(df["Nama Lembaga"].nunique() if "Nama Lembaga" in df.columns else 0)),
        ])
        story.append(Paragraph("Analisis Grafik", styles["PdfHead"]))
        charts = _matplotlib_training_charts(df)
        for i, chart in enumerate(charts):
            story.append(Image(chart, width=16.5*cm, height=8.0*cm))
            if i < len(charts)-1: story.append(Spacer(1, 5))
        story.append(Paragraph("Insight Otomatis", styles["PdfHead"]))
        for text in insights_training(df):
            story.append(Paragraph("• " + _pdf_text(text), styles["PdfBody"]))
        story.append(Paragraph("Kesimpulan", styles["PdfHead"]))
        story.append(Paragraph("Analisis pelatihan menggambarkan jumlah peserta, program, pemerataan provinsi dan lembaga, serta pola kelulusan dan metode pelatihan berdasarkan filter yang dipilih.", styles["PdfBody"]))

    if rekap_df is not None and not rekap_df.empty:
        story.append(PageBreak())
        story.append(Paragraph("🏢 Status Realisasi Lembaga", styles["PdfHead"]))
        sudah = rekap_df[rekap_df["Status Realisasi"] == "Sudah Terealisasi"]
        belum = rekap_df[rekap_df["Status Realisasi"] == "Belum Terealisasi"]
        pct_sudah = (len(sudah) / len(rekap_df) * 100) if len(rekap_df) else 0
        add_kpis([
            ("Sudah Terealisasi", format_number(len(sudah))),
            ("Belum Terealisasi", format_number(len(belum))),
            ("Persentase Terealisasi", f"{pct_sudah:.1f}%")
        ])
        table_data = [["Lembaga/Instansi", "Target", "Realisasi", "Status"]]
        for _, row in rekap_df.sort_values(["Status Realisasi", "Lembaga/Instansi"]).iterrows():
            table_data.append([
                _pdf_text(row.get("Lembaga/Instansi", "")),
                format_number(row.get("Target Pelatihan P3", 0)),
                format_number(row.get("Realisasi Pelatihan P3", 0)),
                _pdf_text(row.get("Status Realisasi", ""))
            ])
        table = Table(table_data, colWidths=[8.0*cm, 2.0*cm, 2.0*cm, 4.0*cm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#EAF2F8")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.HexColor("#17365D")),
            ("GRID", (0,0), (-1,-1), 0.4, colors.lightgrey),
            ("FONTSIZE", (0,0), (-1,-1), 7.5),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("TOPPADDING", (0,0), (-1,-1), 5),
            ("BOTTOMPADDING", (0,0), (-1,-1), 5),
        ]))
        story.append(table)

    if rekap_bim_df is not None and not rekap_bim_df.empty:
        # Status realisasi Bimbingan dibuat konsisten dengan Pelatihan Produktivitas.
        story.append(Paragraph("🏢 Status Realisasi Bimbingan Konsultasi", styles["PdfHead"]))
        sudah = rekap_bim_df[rekap_bim_df["Status Realisasi"] == "Sudah Terealisasi"]
        belum = rekap_bim_df[rekap_bim_df["Status Realisasi"] == "Belum Terealisasi"]
        pct_sudah = (len(sudah) / len(rekap_bim_df) * 100) if len(rekap_bim_df) else 0
        add_kpis([
            ("Sudah Terealisasi", format_number(len(sudah))),
            ("Belum Terealisasi", format_number(len(belum))),
            ("Persentase Terealisasi", f"{pct_sudah:.1f}%")
        ])

        table_data = [["Lembaga/Instansi", "Target", "Realisasi", "Status"]]
        for _, row in rekap_bim_df.sort_values(["Status Realisasi", "Lembaga/Instansi"]).iterrows():
            table_data.append([
                _pdf_text(row.get("Lembaga/Instansi", "")),
                format_number(row.get("Target Perusahaan", 0)),
                format_number(row.get("Realisasi Bimbingan", 0)),
                _pdf_text(row.get("Status Realisasi", ""))
            ])
        table = Table(table_data, colWidths=[8.0*cm, 2.0*cm, 2.0*cm, 4.0*cm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#EAF2F8")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.HexColor("#17365D")),
            ("GRID", (0,0), (-1,-1), 0.4, colors.lightgrey),
            ("FONTSIZE", (0,0), (-1,-1), 7.5),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("TOPPADDING", (0,0), (-1,-1), 5),
            ("BOTTOMPADDING", (0,0), (-1,-1), 5),
        ]))
        story.append(table)

    if bim_df is not None and not bim_df.empty:
        if training_df is not None and not training_df.empty:
            story.append(PageBreak())
        df = bim_df
        story.append(Paragraph("🤝 Bimbingan Konsultasi", styles["PdfHead"]))
        add_kpis([
            ("Total Kegiatan", format_number(len(df))),
            ("Total Perusahaan", format_number(df["NAMA PERUSAHAAN"].nunique() if "NAMA PERUSAHAAN" in df.columns else 0)),
            ("Total Wilayah", format_number(df["NAMA KABUPATEN/KOTA"].nunique() if "NAMA KABUPATEN/KOTA" in df.columns else 0)),
            ("Bidang Usaha", format_number(df["Bidang Usaha Kategori"].nunique() if "Bidang Usaha Kategori" in df.columns else 0)),
        ])
        story.append(Paragraph("Analisis Grafik", styles["PdfHead"]))
        charts = _matplotlib_bim_charts(df)
        for i, chart in enumerate(charts):
            story.append(Image(chart, width=16.5*cm, height=8.0*cm))
            if i < len(charts)-1: story.append(Spacer(1, 5))
        story.append(Paragraph("Insight Otomatis", styles["PdfHead"]))
        for text in insights_bim(df):
            story.append(Paragraph("• " + _pdf_text(text), styles["PdfBody"]))
        story.append(Paragraph("Kesimpulan", styles["PdfHead"]))
        story.append(Paragraph("Analisis bimbingan memberikan gambaran mengenai konsentrasi kegiatan berdasarkan bidang usaha, wilayah, dan waktu pelaksanaan untuk mendukung monitoring dan evaluasi pendampingan.", styles["PdfBody"]))

    if not story:
        story.append(Paragraph("Tidak ada data sesuai filter.", styles["PdfBody"]))
    doc.build(story)
    output.seek(0)
    return output.getvalue()


# =========================================================
# EXPORT HASIL ANALISIS
# =========================================================

def create_analysis_excel(training_df=None, bim_df=None, rekap_bim_df=None, rekap_training_df=None):
    """Membuat workbook Excel analisis yang rapi, terformat, dan siap dibagikan."""
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    output = BytesIO()
    NAVY, TEAL, LIGHT, LIGHER, LINE, WHITE, DARK = "12324A", "0B8F91", "EAF5F7", "F6FAFB", "C7DCE5", "FFFFFF", "19354A"

    def write_df(writer, df, sheet_name, title=None, make_table=True):
        if df is None or df.empty:
            return
        name = sheet_name[:31]
        ws = writer.book.create_sheet(name)
        data = df.copy()

        # Excel Table mensyaratkan nama kolom unik dan tidak boleh kosong.
        # Beberapa file sumber dapat memiliki header yang sama, sehingga
        # kolom dinormalisasi terlebih dahulu agar proses export tidak gagal.
        original_columns = [str(c).strip() if str(c).strip() else "Kolom" for c in data.columns]
        seen = {}
        unique_columns = []
        for col in original_columns:
            count = seen.get(col, 0) + 1
            seen[col] = count
            unique_columns.append(col if count == 1 else f"{col} ({count})")
        data.columns = unique_columns

        # Ubah object kompleks menjadi teks agar aman ditulis ke XLSX.
        for col in data.columns:
            if data[col].dtype == "object":
                data[col] = data[col].map(
                    lambda v: str(v) if isinstance(v, (list, tuple, set, dict)) else v
                )
        if title:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(data.columns)))
            c = ws.cell(1, 1, title)
            c.font = Font(name="Aptos Display", size=14, bold=True, color=WHITE)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(vertical="center")
            ws.row_dimensions[1].height = 27
            header_row = 3
        else:
            header_row = 1
        for j, col in enumerate(data.columns, 1):
            c = ws.cell(header_row, j, str(col))
            c.font = Font(name="Aptos", bold=True, color=WHITE)
            c.fill = PatternFill("solid", fgColor=TEAL)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(bottom=Side(style="thin", color=LINE))
        for i, row in enumerate(data.itertuples(index=False, name=None), header_row + 1):
            for j, value in enumerate(row, 1):
                c = ws.cell(i, j, value)
                c.font = Font(name="Aptos", size=10, color=DARK)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if i % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=LIGHER)
        last_col = get_column_letter(len(data.columns))
        last_row = header_row + len(data)
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
        ws.sheet_view.showGridLines = False
        for j, col in enumerate(data.columns, 1):
            vals = [str(col)] + [str(v) if v is not None else "" for v in data.iloc[:, j-1].head(150)]
            ws.column_dimensions[get_column_letter(j)].width = min(max(max(map(len, vals)) + 2, 11), 36)
        if make_table and len(data) > 0:
            safe = "Tbl" + "".join(ch for ch in sheet_name if ch.isalnum())[:20]
            tab = Table(displayName=safe, ref=f"A{header_row}:{last_col}{last_row}")
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(tab)

    def summary(writer, sheet, title, items):
        write_df(writer, pd.DataFrame(items, columns=["Indikator", "Nilai"]), sheet, title, make_table=False)
        ws = writer.book[sheet]
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 22
        for r in range(4, 4 + len(items)):
            ws.cell(r, 2).font = Font(name="Aptos", size=11, bold=True, color=TEAL)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        info = pd.DataFrame([
            ["Laporan", "Hasil Analisis Dashboard Produktivitas 2026"],
            ["Sumber", "Data sesuai filter aktif pada dashboard"],
            ["Isi", "Data detail, KPI, distribusi, dan insight analisis"],
            ["Catatan", "Setiap sheet dapat difilter dan di-sort untuk eksplorasi lebih lanjut."],
        ], columns=["Keterangan", "Detail"])
        write_df(writer, info, "Petunjuk", "Hasil Analisis Dashboard Produktivitas 2026", make_table=False)
        ws = writer.book["Petunjuk"]; ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 78

        if training_df is not None and not training_df.empty:
            write_df(writer, training_df, "Data Pelatihan", "Data Pelatihan — Sesuai Filter Aktif")
            summary(writer, "KPI Pelatihan", "KPI Pelatihan Produktivitas", [
                ["Total Peserta", len(training_df)],
                ["Total Program", training_df["Judul Program Pelatihan"].nunique() if "Judul Program Pelatihan" in training_df.columns else 0],
                ["Total Provinsi", training_df["Provinsi"].nunique() if "Provinsi" in training_df.columns else 0],
                ["Total Lembaga", training_df["Nama Lembaga"].nunique() if "Nama Lembaga" in training_df.columns else 0],
            ])
            write_df(writer, monthly_data(training_df, "Bulan Pelatihan", "Jumlah Peserta"), "Peserta per Bulan", "Distribusi Peserta per Bulan")
            if "Provinsi" in training_df.columns:
                x = training_df["Provinsi"].fillna("Tidak Diisi").value_counts().reset_index(); x.columns = ["Provinsi", "Jumlah Peserta"]
                write_df(writer, x, "Peserta per Provinsi", "Distribusi Peserta per Provinsi")
            if "Status KeLulusan" in training_df.columns:
                x = training_df["Status KeLulusan"].fillna("Tidak Diisi").value_counts().reset_index(); x.columns = ["Status Kelulusan", "Jumlah"]
                write_df(writer, x, "Status Kelulusan", "Distribusi Status Kelulusan")
            if "Metode Pelatihan" in training_df.columns:
                x = training_df["Metode Pelatihan"].fillna("Tidak Diisi").value_counts().reset_index(); x.columns = ["Metode Pelatihan", "Jumlah"]
                write_df(writer, x, "Metode Pelatihan", "Distribusi Metode Pelatihan")
            write_df(writer, pd.DataFrame({"Insight": [x.replace("**", "") for x in insights_training(training_df)]}), "Insight Pelatihan", "Insight Analisis Pelatihan", False)

        if bim_df is not None and not bim_df.empty:
            write_df(writer, bim_df, "Data Bimbingan", "Data Bimbingan — Sesuai Filter Aktif")

            # Rekap status realisasi lembaga/instansi ikut dimasukkan agar
            # hasil Excel Bimbingan Konsultasi selengkap tampilan dashboard.
            if rekap_bim_df is not None and not rekap_bim_df.empty:
                write_df(
                    writer,
                    rekap_bim_df,
                    "Status Realisasi",
                    "Status Realisasi Bimbingan Konsultasi"
                )

                status_counts = (
                    rekap_bim_df["Status Realisasi"]
                    .fillna("Tidak Diisi")
                    .value_counts()
                    .reindex(["Sudah Terealisasi", "Belum Terealisasi"], fill_value=0)
                    .reset_index()
                )
                status_counts.columns = ["Status Realisasi", "Jumlah Lembaga"]
                write_df(
                    writer,
                    status_counts,
                    "Ringkasan Status",
                    "Ringkasan Status Realisasi",
                    False
                )

            summary(writer, "KPI Bimbingan", "KPI Bimbingan Konsultasi", [
                ["Total Kegiatan", len(bim_df)],
                ["Total Perusahaan", bim_df["NAMA PERUSAHAAN"].nunique() if "NAMA PERUSAHAAN" in bim_df.columns else 0],
                ["Total Wilayah", bim_df["NAMA KABUPATEN/KOTA"].nunique() if "NAMA KABUPATEN/KOTA" in bim_df.columns else 0],
                ["Kategori Bidang Usaha", bim_df["Bidang Usaha Kategori"].nunique() if "Bidang Usaha Kategori" in bim_df.columns else 0],
            ])

            # Rekap analisis yang sama dengan grafik dashboard.
            if "Bulan Bimbingan" in bim_df.columns:
                write_df(
                    writer,
                    monthly_data(bim_df, "Bulan Bimbingan", "Jumlah Kegiatan"),
                    "Bimbingan per Bulan",
                    "Distribusi Bimbingan per Bulan"
                )
            if "Bidang Usaha Kategori" in bim_df.columns:
                x = bim_df["Bidang Usaha Kategori"].fillna("Tidak Diisi").value_counts().reset_index()
                x.columns = ["Bidang Usaha", "Jumlah Kegiatan"]
                write_df(writer, x, "Bidang Usaha", "Distribusi Bimbingan per Bidang Usaha")
            if "NAMA KABUPATEN/KOTA" in bim_df.columns:
                x = bim_df["NAMA KABUPATEN/KOTA"].fillna("Tidak Diisi").value_counts().reset_index()
                x.columns = ["Wilayah", "Jumlah Kegiatan"]
                write_df(writer, x, "Wilayah", "Distribusi Bimbingan per Wilayah")

            write_df(
                writer,
                pd.DataFrame({"Insight": [x.replace("**", "") for x in insights_bim(bim_df)]}),
                "Insight Bimbingan",
                "Insight Analisis Bimbingan",
                False
            )

    output.seek(0)
    return output.getvalue()


# =========================================================
# HALAMAN AWAL / PEMILIHAN FILE
# =========================================================
# Halaman pembuka hanya ditampilkan sebelum file Excel dipilih.
# Setelah file dipilih, aplikasi langsung masuk ke menu dashboard.

if "excel_files" not in st.session_state:
    st.session_state.excel_files = None

if "excel_bytes" not in st.session_state:
    st.session_state.excel_bytes = None
    st.session_state.excel_name = ""

if st.session_state.excel_files is None:
    st.markdown("### MONITORING PROGRAM • 2026")

    hero_left, hero_right = st.columns([1.35, 0.65], gap="large", vertical_alignment="center")

    with hero_left:
        st.markdown("# Dashboard Produktivitas 2026")
        st.markdown("### Pelatihan Produktivitas & Bimbingan Konsultasi")
        st.write(
            "Platform monitoring dan analisis untuk melihat capaian program, "
            "sebaran kegiatan, realisasi lembaga, serta temuan utama sebagai "
            "bahan evaluasi dan pengambilan keputusan."
        )


    st.divider()

    st.markdown("### Fokus Dashboard")
    c1, c2, c3, c4 = st.columns(4, gap="medium")

    with c1:
        with st.container(border=True):
            st.markdown("#### 🎓 Pelatihan")
            st.caption("Peserta, program, provinsi, metode, dan kelulusan.")

    with c2:
        with st.container(border=True):
            st.markdown("#### 🤝 Bimbingan")
            st.caption("Perusahaan, wilayah, bidang usaha, dan kegiatan.")

    with c3:
        with st.container(border=True):
            st.markdown("#### 🏢 Realisasi")
            st.caption("Lembaga yang sudah dan belum terealisasi.")

    with c4:
        with st.container(border=True):
            st.markdown("#### 💡 Insight")
            st.caption("Temuan utama, tren dan ringkasan analisis.")

    st.divider()

    st.markdown("### Mulai Analisis Data")
    st.caption(
        "Upload file Excel untuk membuka dashboard. Data akan dibersihkan dan "
        "distandarkan secara otomatis."
    )

    upload_box = st.container(border=True)
    with upload_box:
        uploaded = st.file_uploader(
            "Upload File Excel",
            type=["xlsx"],
            accept_multiple_files=True,
            help="Upload satu atau beberapa file Excel. Sheet dengan nama yang sama akan digabung otomatis.",
            key="excel_uploader"
        )

    if not uploaded:
        st.info("Pilih satu atau beberapa file Excel untuk melanjutkan ke dashboard analisis.")
        st.stop()

    st.session_state.excel_files = [
        (file.name, file.getvalue())
        for file in uploaded
    ]
    st.session_state.excel_bytes = st.session_state.excel_files[0][1]
    st.session_state.excel_name = st.session_state.excel_files[0][0]
    st.rerun()


# File yang sudah dipilih disimpan di session agar halaman pembuka tidak
# muncul lagi ketika pengguna berpindah menu.
file_items = st.session_state.excel_files or [(st.session_state.excel_name, st.session_state.excel_bytes)]
file_bytes = file_items[0][1]

# =========================================================
# PROSES DATA
# =========================================================

with st.spinner(
    "⏳ Membaca dan membersihkan data..."
):

    sheets = load_excel(
        tuple(file_items)
    )


    train = (

        clean_training(
            sheets[TRAIN_SHEET]
        )

        if TRAIN_SHEET in sheets

        else pd.DataFrame()
    )


    bim = (

        clean_bim(
            sheets[BIM_SHEET]
        )

        if BIM_SHEET in sheets

        else pd.DataFrame()
    )


    rekap = (

        clean_rekap_realisasi(
            sheets[REKAP_SHEET]
        )

        if REKAP_SHEET in sheets

        else pd.DataFrame()
    )


    rekap_bim = (

        clean_rekap_bimkon(
            sheets[REKAP_BIM_SHEET]
        )

        if REKAP_BIM_SHEET in sheets

        else pd.DataFrame()
    )


st.markdown(
    f"""
    <div class="data-note">
        <span>✓</span>
        <span><strong>Data siap dianalisis.</strong> {format_number(len(train))} data pelatihan dan {format_number(len(bim))} data bimbingan berhasil diproses.</span>
    </div>
    """,
    unsafe_allow_html=True
)


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.markdown(
    """
    <div class="sidebar-brand">
        <div class="brand-mark">◈</div>
        <div class="brand-title">Produktivitas</div>
        <div class="brand-sub">Monitoring Dashboard • 2026</div>
    </div>
    """,
    unsafe_allow_html=True
)


page = st.sidebar.radio(
    "Menu",
    [
        "🏠 Overview",
        "🎓 Pelatihan Produktivitas",
        "🤝 Bimbingan Konsultasi",
        "📋 Data Detail"
    ]
)


st.sidebar.divider()

# File Excel hanya dipilih pada halaman awal agar sidebar tetap bersih.
# Tidak ada uploader tambahan di atas filter.


# =========================================================
# OVERVIEW
# =========================================================

if page == "🏠 Overview":

    page_hero(
        "MONITORING PROGRAM • 2026",
        "Dashboard Produktivitas 2026",
        "Ringkasan pelaksanaan Pelatihan Produktivitas dan Bimbingan Konsultasi untuk monitoring, evaluasi, dan pelaporan.",
        show_decorative=True
    )


    total_program = (

        train[
            "Judul Program Pelatihan"
        ].nunique()

        if "Judul Program Pelatihan"
        in train.columns

        else 0
    )


    total_perusahaan = (

        bim[
            "NAMA PERUSAHAAN"
        ].nunique()

        if "NAMA PERUSAHAAN"
        in bim.columns

        else 0
    )


    a, b, c, d = st.columns(4)


    with a:

        st.metric(
            "👥 Total Peserta",
            format_number(
                len(train)
            )
        )


    with b:

        st.metric(
            "📚 Total Program",
            format_number(
                total_program
            )
        )

    with c:

        st.metric(
            "🤝 Total Bimbingan",
            format_number(
                len(bim)
            )
        )


    with d:

        st.metric(
            "🏢 Total Perusahaan",
            format_number(
                total_perusahaan
            )
        )



    section_heading("📊 Gambaran Umum Data", "Ringkasan capaian Pelatihan Produktivitas dan Bimbingan Konsultasi berdasarkan data yang tersedia pada tahun 2026.")


    col1, col2 = st.columns(2)


    # -----------------------------------------------------
    # PESERTA
    # -----------------------------------------------------

    with col1:

        x = monthly_data(
            train,
            "Bulan Pelatihan",
            "Peserta"
        )

        fig = px.line(
            x,
            x="Bulan",
            y="Peserta",
            markers=True,
            title="🎓 Perkembangan Peserta Pelatihan"
        )

        fig.update_traces(
            line_width=3, marker_size=7,
            hovertemplate="<b>%{x}</b><br>Peserta: %{y}<extra></extra>"
        )
        fig.update_xaxes(
            tickmode="array", tickvals=MONTHS,
            ticktext=short_month_labels(), tickangle=0, title_text=None
        )
        fig.update_yaxes(title_text=None, rangemode="tozero", tickformat=",d")

        st.plotly_chart(
            style_chart(fig, 360),
            use_container_width=True
        )


    # -----------------------------------------------------
    # BIMBINGAN
    # -----------------------------------------------------

    with col2:

        x = monthly_data(
            bim,
            "Bulan Bimbingan",
            "Kegiatan"
        )

        fig = px.bar(
            x,
            x="Bulan",
            y="Kegiatan",
            title="🤝 Perkembangan Bimbingan Konsultasi"
        )

        fig.update_traces(
            cliponaxis=False,
            hovertemplate="<b>%{x}</b><br>Kegiatan: %{y}<extra></extra>"
        )
        fig.update_xaxes(
            tickmode="array",
            tickvals=MONTHS,
            ticktext=short_month_labels(),
            tickangle=0,
            title_text=None
        )
        fig.update_yaxes(title_text=None, rangemode="tozero", tickformat=",d")

        st.plotly_chart(
            style_chart(fig, 360),
            use_container_width=True
        )


    # -----------------------------------------------------
    # KESIMPULAN
    # -----------------------------------------------------

    st.subheader("🎯 Kesimpulan Analisis")
    st.info(
        "Dashboard memberikan gambaran mengenai jumlah peserta, program pelatihan, "
        "wilayah, kegiatan bimbingan, perusahaan, bidang usaha, serta perkembangan "
        "aktivitas sepanjang periode data yang tersedia.\n\n"
        "Informasi tersebut dapat digunakan sebagai dasar monitoring, evaluasi, "
        "dan penyusunan laporan pelaksanaan program produktivitas."
    )


# =========================================================
# PELATIHAN
# =========================================================

elif page == "🎓 Pelatihan Produktivitas":

    page_hero(
        "PELATIHAN",
        "Pelatihan Produktivitas",
        "Pantau peserta, program, sebaran provinsi, metode pelatihan, dan status kelulusan secara ringkas."
    )


    st.sidebar.markdown(
        "### 🔎 Filter Pelatihan"
    )


    f = train.copy()


    # Filter

    for col, label, key in [

        (
            "Tahun Pelatihan",
            "Tahun",
            "train_year"
        ),

        (
            "Provinsi",
            "Provinsi",
            "train_province"
        ),

        (
            "Metode Pelatihan",
            "Metode",
            "train_method"
        ),

        (
            "Jenis Pelatihan",
            "Jenis Pelatihan",
            "train_type"
        ),

        (
            "Status KeLulusan",
            "Status Kelulusan",
            "train_status"
        )

    ]:

        f = selectbox_filter(
            f,
            col,
            label,
            key
        )


    st.caption(
        f"📌 Menampilkan **{format_number(len(f))} peserta** "
        "sesuai filter."
    )


    # -----------------------------------------------------
    # KPI
    # -----------------------------------------------------

    a, b, c, d = st.columns(4)


    with a:

        st.metric(
            "👥 Total Peserta",
            format_number(
                len(f)
            )
        )


    with b:

        total_program = (

            f[
                "Judul Program Pelatihan"
            ].nunique()

            if "Judul Program Pelatihan"
            in f.columns

            else 0
        )

        st.metric(
            "📚 Total Program",
            format_number(
                total_program
            )
        )

    with c:

        total_provinsi = (

            f[
                "Provinsi"
            ].nunique()

            if "Provinsi"
            in f.columns

            else 0
        )

        st.metric(
            "📍 Total Provinsi",
            format_number(
                total_provinsi
            )
        )


    with d:

        total_lembaga = (

            f[
                "Nama Lembaga"
            ].nunique()

            if "Nama Lembaga"
            in f.columns

            else 0
        )

        st.metric(
            "🏢 Total Lembaga",
            format_number(
                total_lembaga
            )
        )


    if "Judul Program Pelatihan" in f.columns:
        show_kpi_tags(
            f["Judul Program Pelatihan"].dropna().unique(),
            "Program pelatihan yang tercakup",
            icon="📚"
        )


    # -----------------------------------------------------
    # STATUS REALISASI LEMBAGA
    # -----------------------------------------------------

    section_heading("🏢 Status Realisasi Lembaga", "Perbandingan lembaga yang sudah dan belum terealisasi.")
    st.caption(
        "Menunjukkan lembaga/instansi yang sudah memiliki realisasi pelatihan P3 dan yang belum terealisasi."
    )

    if not rekap.empty:
        sudah = rekap[rekap["Status Realisasi"] == "Sudah Terealisasi"].copy()
        belum = rekap[rekap["Status Realisasi"] == "Belum Terealisasi"].copy()
        total_lembaga_rekap = len(rekap)
        pct_sudah = (len(sudah) / total_lembaga_rekap * 100) if total_lembaga_rekap else 0

        r1, r2, r3 = st.columns(3)
        with r1:
            st.metric("Sudah Terealisasi", format_number(len(sudah)))
        with r2:
            st.metric("Belum Terealisasi", format_number(len(belum)))
        with r3:
            st.metric("Persentase Lembaga Terealisasi", f"{pct_sudah:.1f}%")

        status_chart = pd.DataFrame({
            "Status": ["Sudah Terealisasi", "Belum Terealisasi"],
            "Jumlah Lembaga": [len(sudah), len(belum)]
        })
        fig_status = px.bar(
            status_chart, x="Status", y="Jumlah Lembaga",
            title="Status Realisasi Lembaga/Instansi"
        )
        fig_status.update_traces(
            hovertemplate="<b>%{x}</b><br>Jumlah lembaga: %{y}<extra></extra>"
        )
        fig_status.update_xaxes(title_text=None, tickangle=0)
        fig_status.update_yaxes(title_text=None, rangemode="tozero", tickformat=",d")
        st.plotly_chart(style_chart(fig_status, 330), use_container_width=True)

        t1, t2 = st.columns(2)
        with t1:
            st.markdown("#### ✅ Sudah Terealisasi")
            cols = [c for c in ["Lembaga/Instansi", "Target Pelatihan P3", "Realisasi Pelatihan P3", "Persentase Realisasi"] if c in sudah.columns]
            st.dataframe(sudah[cols].sort_values("Realisasi Pelatihan P3", ascending=False), use_container_width=True, hide_index=True)

        with t2:
            st.markdown("#### ⏳ Belum Terealisasi")
            cols = [c for c in ["Lembaga/Instansi", "Target Pelatihan P3", "Realisasi Pelatihan P3", "Persentase Realisasi"] if c in belum.columns]
            st.dataframe(belum[cols].sort_values("Lembaga/Instansi"), use_container_width=True, hide_index=True)
    else:
        st.info(
            "Data rekap realisasi lembaga belum tersedia pada file Excel. Pastikan terdapat sheet **Rekap Pel Prod** dengan kolom Lembaga/Instansi, Target, dan Realisasi."
        )


    # -----------------------------------------------------
    # ANALISIS
    # -----------------------------------------------------

    section_heading("📊 Analisis Pelatihan", "Distribusi peserta dan karakteristik pelatihan.")


    col1, col2 = st.columns(2)


    # Tren peserta

    with col1:

        x = monthly_data(
            f,
            "Bulan Pelatihan",
            "Peserta"
        )

        fig = px.line(
            x,
            x="Bulan",
            y="Peserta",
            markers=True,
            title="📈 Peserta per Bulan"
        )

        fig.update_traces(
            mode="lines+markers",
            line_width=3,
            marker_size=7,
            hovertemplate="<b>%{x}</b><br>Peserta: %{y}<extra></extra>"
        )
        fig.update_xaxes(
            tickmode="array",
            tickvals=MONTHS,
            ticktext=short_month_labels(),
            tickangle=0,
            title_text=None
        )
        fig.update_yaxes(title_text=None, rangemode="tozero", tickformat=",d")

        st.plotly_chart(
            style_chart(fig, 320),
            use_container_width=True
        )


    # Top provinsi

    with col2:

        if "Provinsi" in f.columns:

            x = (
                f["Provinsi"]
                .fillna("Tidak Diisi")
                .value_counts()
                .head(10)
                .sort_values()
                .reset_index()
            )

            x.columns = [
                "Provinsi",
                "Peserta"
            ]

            fig = px.bar(
                x,
                x="Peserta",
                y="Provinsi",
                orientation="h",
                title="🏆 Top 10 Provinsi"
            )

            fig.update_traces(
                hovertemplate="<b>%{y}</b><br>Peserta: %{x}<extra></extra>"
            )
            fig.update_xaxes(title_text=None, rangemode="tozero", tickformat=",d")
            fig.update_yaxes(title_text=None, autorange=True)

            st.plotly_chart(
                style_chart(fig, 340, horizontal=True),
                use_container_width=True
            )


    # -----------------------------------------------------
    # STATUS + METODE
    # -----------------------------------------------------

    col1, col2 = st.columns(2)


    with col1:

        if "Status KeLulusan" in f.columns:

            x = (
                f[
                    "Status KeLulusan"
                ]
                .fillna("Tidak Diisi")
                .value_counts()
                .reset_index()
            )

            x.columns = [
                "Status",
                "Jumlah"
            ]

            fig = px.pie(
                x,
                names="Status",
                values="Jumlah",
                hole=0.62,
                title="🎓 Status Kelulusan"
            )

            fig.update_traces(
                textinfo="none",
                hovertemplate="<b>%{label}</b><br>Jumlah: %{value}<br>Proporsi: %{percent}<extra></extra>"
            )
            fig.update_layout(
                showlegend=True,
                legend=dict(
                    orientation="h",
                    yanchor="bottom",
                    y=-0.05,
                    x=0.5,
                    xanchor="center",
                    font=dict(size=10)
                )
            )

            st.plotly_chart(
                style_chart(fig, 340, show_legend=True),
                use_container_width=True
            )


    with col2:

        if "Metode Pelatihan" in f.columns:

            x = (
                f[
                    "Metode Pelatihan"
                ]
                .fillna("Tidak Diisi")
                .value_counts()
                .reset_index()
            )

            x.columns = [
                "Metode",
                "Jumlah"
            ]

            fig = px.bar(
                x,
                x="Metode",
                y="Jumlah",
                title="⚙️ Metode Pelatihan"
            )

            fig.update_traces(
                hovertemplate="<b>%{x}</b><br>Peserta: %{y}<extra></extra>"
            )
            fig.update_xaxes(title_text=None, tickangle=0, automargin=True)
            fig.update_yaxes(title_text=None, rangemode="tozero", tickformat=",d")

            st.plotly_chart(
                style_chart(fig, 320),
                use_container_width=True
            )


    # -----------------------------------------------------
    # INSIGHT
    # -----------------------------------------------------

    section_heading("💡 Insight Otomatis", "Temuan yang dihasilkan otomatis dari data terfilter.")


    if f.empty:

        st.warning(
            "Tidak ada data sesuai filter."
        )

    else:

        for text in insights_training(f):
            insight_card(text)


    # -----------------------------------------------------
    # DOWNLOAD HASIL ANALISIS
    # -----------------------------------------------------

    st.subheader("📥 Download Hasil Analisis")
    st.caption(
        "Download PDF berisi hasil akhir dashboard sesuai filter aktif: KPI, grafik, "
        "insight, dan kesimpulan analisis."
    )

    if not f.empty:
        pdf_training = create_dashboard_pdf(training_df=f, rekap_df=rekap)
        st.download_button(
            label="📄 Download Hasil Dashboard Pelatihan (PDF)",
            data=pdf_training,
            file_name="Hasil_Dashboard_Pelatihan_Produktivitas_2026.pdf",
            mime="application/pdf",
            use_container_width=True
        )


# =========================================================
# BIMBINGAN
# =========================================================

elif page == "🤝 Bimbingan Konsultasi":

    page_hero(
        "BIMBINGAN & KONSULTASI",
        "Bimbingan Konsultasi",
        "Pantau kegiatan pendampingan berdasarkan perusahaan, wilayah, bidang usaha, dan pelaksana."
    )


    st.sidebar.markdown(
        "### 🔎 Filter Bimbingan"
    )


    f = bim.copy()


    for col, label, key in [

        (
            "Bidang Usaha Kategori",
            "Bidang Usaha",
            "bim_business"
        ),

        (
            "NAMA KABUPATEN/KOTA",
            "Kabupaten/Kota",
            "bim_region"
        ),

        (
            "PELAKSANA",
            "Pelaksana",
            "bim_executor"
        ),

        (
            "Tahun Bimbingan",
            "Tahun",
            "bim_year"
        )

    ]:

        f = selectbox_filter(
            f,
            col,
            label,
            key
        )


    st.caption(
        f"📌 Menampilkan **{format_number(len(f))} kegiatan** "
        "sesuai filter."
    )


    # -----------------------------------------------------
    # KPI
    # -----------------------------------------------------

    a, b, c, d = st.columns(4)


    with a:

        st.metric(
            "📋 Total Kegiatan",
            format_number(
                len(f)
            )
        )


    with b:

        total_perusahaan = (

            f[
                "NAMA PERUSAHAAN"
            ].nunique()

            if "NAMA PERUSAHAAN"
            in f.columns

            else 0
        )

        st.metric(
            "🏢 Total Perusahaan",
            format_number(
                total_perusahaan
            )
        )


    with c:

        total_wilayah = (

            f[
                "NAMA KABUPATEN/KOTA"
            ].nunique()

            if "NAMA KABUPATEN/KOTA"
            in f.columns

            else 0
        )

        st.metric(
            "📍 Total Wilayah",
            format_number(
                total_wilayah
            )
        )


    with d:

        total_bidang = (

            f[
                "Bidang Usaha Kategori"
            ].nunique()

            if "Bidang Usaha Kategori"
            in f.columns

            else 0
        )

        st.metric(
            "🏭 Bidang Usaha",
            format_number(
                total_bidang
            )
        )


    # -----------------------------------------------------
    # STATUS REALISASI LEMBAGA / INSTANSI
    # -----------------------------------------------------

    section_heading("🏢 Status Realisasi Bimbingan Konsultasi", "Perbandingan lembaga/instansi yang sudah dan belum terealisasi.")
    st.caption(
        "Menunjukkan lembaga/instansi yang sudah memiliki realisasi bimbingan konsultasi dan yang belum terealisasi."
    )

    if not rekap_bim.empty:
        sudah_bim = rekap_bim[
            rekap_bim["Status Realisasi"] == "Sudah Terealisasi"
        ].copy()
        belum_bim = rekap_bim[
            rekap_bim["Status Realisasi"] == "Belum Terealisasi"
        ].copy()
        total_rekap_bim = len(rekap_bim)
        pct_sudah_bim = (
            len(sudah_bim) / total_rekap_bim * 100
            if total_rekap_bim else 0
        )

        r1, r2, r3 = st.columns(3)
        with r1:
            st.metric("Sudah Terealisasi", format_number(len(sudah_bim)))
        with r2:
            st.metric("Belum Terealisasi", format_number(len(belum_bim)))
        with r3:
            st.metric(
                "Persentase Lembaga Terealisasi",
                f"{pct_sudah_bim:.1f}%"
            )

        status_bim_chart = pd.DataFrame({
            "Status": ["Sudah Terealisasi", "Belum Terealisasi"],
            "Jumlah Lembaga": [len(sudah_bim), len(belum_bim)]
        })
        fig_status_bim = px.bar(
            status_bim_chart,
            x="Status",
            y="Jumlah Lembaga",
            title="Status Realisasi Bimbingan Konsultasi"
        )
        fig_status_bim.update_traces(
            hovertemplate="<b>%{x}</b><br>Jumlah lembaga: %{y}<extra></extra>"
        )
        fig_status_bim.update_xaxes(title_text=None, tickangle=0)
        fig_status_bim.update_yaxes(title_text=None, rangemode="tozero", tickformat=",d")
        st.plotly_chart(
            style_chart(fig_status_bim, 330),
            use_container_width=True
        )

        t1, t2 = st.columns(2)
        with t1:
            st.markdown("#### ✅ Sudah Terealisasi")
            cols_bim = [
                c for c in [
                    "Lembaga/Instansi",
                    "Target Perusahaan",
                    "Realisasi Bimbingan",
                    "Persentase Realisasi"
                ] if c in sudah_bim.columns
            ]
            st.dataframe(
                sudah_bim[cols_bim].sort_values(
                    "Realisasi Bimbingan", ascending=False
                ),
                use_container_width=True,
                hide_index=True
            )

        with t2:
            st.markdown("#### ⏳ Belum Terealisasi")
            cols_bim = [
                c for c in [
                    "Lembaga/Instansi",
                    "Target Perusahaan",
                    "Realisasi Bimbingan",
                    "Persentase Realisasi"
                ] if c in belum_bim.columns
            ]
            st.dataframe(
                belum_bim[cols_bim].sort_values(
                    "Lembaga/Instansi"
                ),
                use_container_width=True,
                hide_index=True
            )
    else:
        st.info(
            "Data rekap realisasi bimbingan belum tersedia. Pastikan terdapat sheet "
            "**Rekap Bimkon** dengan kolom Lembaga/Instansi, Target Perusahaan, dan Realisasi."
        )


    # -----------------------------------------------------
    # GRAFIK
    # -----------------------------------------------------

    section_heading("📊 Analisis Bimbingan", "Distribusi kegiatan berdasarkan bidang usaha, wilayah, dan waktu.")


    col1, col2 = st.columns(2)


    # Bidang usaha

    with col1:

        if "Bidang Usaha Kategori" in f.columns:

            x = (
                f[
                    "Bidang Usaha Kategori"
                ]
                .fillna("Tidak Diisi")
                .value_counts()
                .sort_values()
                .reset_index()
            )

            x.columns = [
                "Bidang Usaha",
                "Jumlah"
            ]

            fig = px.bar(
                x,
                x="Jumlah",
                y="Bidang Usaha",
                orientation="h",
                title="🏭 Bimbingan per Bidang Usaha"
            )

            fig.update_traces(
                hovertemplate="<b>%{y}</b><br>Jumlah: %{x}<extra></extra>"
            )
            fig.update_xaxes(title_text=None, rangemode="tozero", tickformat=",d")
            fig.update_yaxes(title_text=None)

            st.plotly_chart(
                style_chart(fig, 340, horizontal=True),
                use_container_width=True
            )


    # Wilayah

    with col2:

        if "NAMA KABUPATEN/KOTA" in f.columns:

            x = (
                f[
                    "NAMA KABUPATEN/KOTA"
                ]
                .fillna("Tidak Diisi")
                .value_counts()
                .head(10)
                .sort_values()
                .reset_index()
            )

            x.columns = [
                "Wilayah",
                "Jumlah"
            ]

            fig = px.bar(
                x,
                x="Jumlah",
                y="Wilayah",
                orientation="h",
                title="📍 Top 10 Wilayah"
            )

            fig.update_traces(
                hovertemplate="<b>%{y}</b><br>Jumlah: %{x}<extra></extra>"
            )
            fig.update_xaxes(title_text=None, rangemode="tozero", tickformat=",d")
            fig.update_yaxes(title_text=None)

            st.plotly_chart(
                style_chart(fig, 340, horizontal=True),
                use_container_width=True
            )


    # -----------------------------------------------------
    # BULAN
    # -----------------------------------------------------

    x = monthly_data(
        f,
        "Bulan Bimbingan",
        "Kegiatan"
    )

    fig = px.line(
        x,
        x="Bulan",
        y="Kegiatan",
        markers=True,
        title="📅 Jumlah Bimbingan per Bulan"
    )

    fig.update_traces(
        line_width=3, marker_size=7,
        hovertemplate="<b>%{x}</b><br>Kegiatan: %{y}<extra></extra>"
    )
    fig.update_xaxes(
        tickmode="array", tickvals=MONTHS,
        ticktext=short_month_labels(), tickangle=0, title_text=None
    )
    fig.update_yaxes(title_text=None, rangemode="tozero", tickformat=",d")

    st.plotly_chart(
        style_chart(fig, 340),
        use_container_width=True
    )


    # -----------------------------------------------------
    # INSIGHT
    # -----------------------------------------------------

    section_heading("💡 Insight Otomatis", "Temuan yang dihasilkan otomatis dari data terfilter.")


    if f.empty:

        st.warning(
            "Tidak ada data sesuai filter."
        )

    else:

        for text in insights_bim(f):
            insight_card(text)


    # -----------------------------------------------------
    # KESIMPULAN
    # -----------------------------------------------------

    st.subheader("🎯 Kesimpulan Analisis Bimbingan")
    st.info(
        "Analisis bimbingan memberikan gambaran mengenai konsentrasi kegiatan "
        "berdasarkan bidang usaha, wilayah, dan waktu pelaksanaan.\n\n"
        "Data tersebut dapat digunakan untuk monitoring pemerataan pendampingan "
        "serta identifikasi wilayah atau bidang usaha yang masih membutuhkan "
        "perhatian lebih lanjut."
    )


    # -----------------------------------------------------
    # DOWNLOAD
    # -----------------------------------------------------

    st.subheader("📥 Download Hasil Analisis")
    st.caption(
        "Download PDF berisi hasil akhir dashboard sesuai filter aktif: KPI, grafik, "
        "insight, dan kesimpulan analisis."
    )

    if not f.empty:
        pdf_bim = create_dashboard_pdf(bim_df=f, rekap_bim_df=rekap_bim)
        st.download_button(
            label="📄 Download Hasil Dashboard Bimbingan (PDF)",
            data=pdf_bim,
            file_name="Hasil_Dashboard_Bimbingan_Konsultasi_2026.pdf",
            mime="application/pdf",
            use_container_width=True
        )


# =========================================================
# DATA DETAIL
# =========================================================

else:

    page_hero(
        "DATA MANAGEMENT",
        "Data Detail & Hasil Cleaning",
        "Periksa data yang telah dibersihkan dan distandarkan sebelum digunakan untuk analisis."
    )


    tipe = st.radio(
        "Pilih data",
        [
            "Pelatihan",
            "Bimbingan"
        ],
        horizontal=True
    )


    df = (
        train
        if tipe == "Pelatihan"
        else bim
    )


    a, b = st.columns(2)


    with a:

        st.metric(
            "📊 Jumlah Baris",
            format_number(
                len(df)
            )
        )


    with b:

        st.metric(
            "🧹 Jumlah Kolom",
            format_number(
                len(df.columns)
            )
        )


    st.dataframe(
        df,
        use_container_width=True,
        height=550
    )


    # =====================================================
    # DOWNLOAD DATA CLEANING
    # =====================================================

    output = BytesIO()


    df.to_excel(
        output,
        index=False,
        engine="openpyxl"
    )


    st.download_button(

        label="⬇️ Download Data Hasil Cleaning",

        data=output.getvalue(),

        file_name=(
            f"data_{tipe.lower()}_clean.xlsx"
        ),

        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        use_container_width=True
    )


# =========================================================
# FOOTER
# =========================================================

